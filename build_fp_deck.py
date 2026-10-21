#!/usr/bin/env python3
"""Deck: closing CodeSecurity false positives across ServiceNow, Azure DevOps, and GitHub.

Reads web/eval/results/results.json (from web/eval/run-avit-eval.cjs) and writes
docs/threatlint-fp-closure.pptx. Run the evaluation first.
"""

import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY   = RGBColor(0x0D, 0x1B, 0x2A)
AMBER  = RGBColor(0xD9, 0x77, 0x06)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHT = RGBColor(0xF4, 0xF6, 0xF9)
GREY   = RGBColor(0x6B, 0x7A, 0x8D)
INK    = RGBColor(0x1A, 0x23, 0x30)
GREEN  = RGBColor(0x1E, 0x8E, 0x4E)
RED    = RGBColor(0xC0, 0x39, 0x2B)
LINE   = RGBColor(0xD5, 0xDB, 0xE3)

ROOT = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(ROOT, "web/eval/results/results.json")
OUT = os.path.join(ROOT, "docs/threatlint-fp-closure.pptx")

W, H = Inches(13.33), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, rgb, line=None):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def txt(slide, text, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), bold, color, font
        para.space_after = Pt(4)
    return box


def light(title, kicker=None):
    slide = prs.slides.add_slide(blank)
    rect(slide, 0, 0, W, H, OFFWHT)
    rect(slide, 0, 0, Inches(0.18), H, AMBER)
    if kicker:
        txt(slide, kicker.upper(), Inches(0.6), Inches(0.35), Inches(12), Inches(0.4), 12, True, AMBER)
    txt(slide, title, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.9), 30, True, NAVY)
    return slide


def card(slide, x, y, w, h, heading, body, accent=AMBER, size=14):
    rect(slide, x, y, w, h, WHITE, LINE)
    rect(slide, x, y, w, Inches(0.07), accent)
    txt(slide, heading, x + Inches(0.2), y + Inches(0.18), w - Inches(0.4), Inches(0.5), 16, True, NAVY)
    txt(slide, body, x + Inches(0.2), y + Inches(0.7), w - Inches(0.4), h - Inches(0.8), size, False, INK)


def table(slide, rows, x, y, widths, header_fill=NAVY, size=11, row_h=0.42):
    shape = slide.shapes.add_table(len(rows), len(widths), x, y, Emu(sum(widths)), Inches(row_h * len(rows)))
    grid = shape.table
    for col, width in enumerate(widths):
        grid.columns[col].width = width
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = grid.cell(r, c)
            para = cell.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = str(value)
            run.font.size, run.font.bold = Pt(size), r == 0
            run.font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else OFFWHT)
    return grid


results = json.load(open(RESULTS, encoding="utf-8"))
m, rows, closures = results["metrics"], results["rows"], results["closures"]
pct = lambda a, b: f"{round(100 * a / b)}%" if b else "n/a"

# 1. Title
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, NAVY)
rect(slide, Inches(0.6), Inches(2.4), Inches(0.12), Inches(2.2), AMBER)
txt(slide, "Closing false positives with evidence", Inches(0.95), Inches(2.3), Inches(11.5), Inches(1), 44, True, WHITE)
txt(slide, "Threatlint Findings triage → ServiceNow AVR · Azure DevOps Boards · GitHub Advanced Security → Semgrep tuning",
    Inches(0.95), Inches(3.4), Inches(11.5), Inches(0.8), 20, False, RGBColor(0xC8, 0xD3, 0xE0))
txt(slide, f"Efficacy measured on {m['total']} CodeSecurity AVITs · {m['model']} · {results['generatedAt'][:10]}",
    Inches(0.95), Inches(4.3), Inches(11.5), Inches(0.5), 16, False, AMBER)

# 2. Problem
slide = light("Scanner noise costs AppSec hours per finding", "The problem")
card(slide, Inches(0.6), Inches(1.8), Inches(3.9), Inches(4.6), "Triage is manual",
     "Every CodeSecurity AVIT needs someone to open the code, trace the data flow, and decide whether it is real.")
card(slide, Inches(4.72), Inches(1.8), Inches(3.9), Inches(4.6), "Closure loses the evidence",
     "Items get closed with a one-line note. Auditors and the next reviewer can't see why it was a false positive.")
card(slide, Inches(8.84), Inches(1.8), Inches(3.9), Inches(4.6), "False positives come back",
     "Without tuning, the same finding reappears on the next scan, on every branch, in every tracker it syncs to.")

# 3. How it works
slide = light("One flow, three trackers", "How it works")
steps = [("1  Fetch", "CodeSecurity items only.\nOther AVR views are skipped and listed as excluded."),
         ("2  Triage", "The FP reviewer reads the repository, traces each finding, and records a disposition with path:line evidence."),
         ("3  Confirm", "A FALSE POSITIVE gets the “Close FP & tune Semgrep…” action. Each step needs explicit confirmation."),
         ("4  Close", "The report is attached first, prefixed with the item ID; the item is then closed with “Closed by Threatlint”."),
         ("5  Tune", "Matching Semgrep findings (same rule, project, file) are ignored as false_positive across scans and branches.")]
for index, (heading, body) in enumerate(steps):
    x = Inches(0.6 + index * 2.5)
    card(slide, x, Inches(1.9), Inches(2.3), Inches(3.6), heading, body, size=13)
    if index < len(steps) - 1:
        txt(slide, "→", x + Inches(2.28), Inches(3.3), Inches(0.3), Inches(0.5), 22, True, AMBER)
txt(slide, "A verdict never closes anything by itself. The server re-checks the disposition, the single tracker reference, and the CodeSecurity scope on every request.",
    Inches(0.6), Inches(5.9), Inches(12.2), Inches(0.8), 14, False, GREY)

# 4. Trackers
slide = light("Same result in each system", "Tracker coverage")
table(slide, [
    ["", "ServiceNow AVR", "Azure DevOps Boards", "GitHub Advanced Security"],
    ["Item reference", "AVIT0012345", "ADO-4821", "GHAS-owner/repo#12"],
    ["Report attached as", "AVIT0012345-threatlint-fp-review.md (+ .docx)", "ADO-4821-threatlint-fp-review.md (+ .docx)", "Tracking issue (alerts cannot hold files)"],
    ["Closing note", "Work note: Closed by Threatlint…", "Discussion comment: Closed by Threatlint…", "Dismissal comment: Closed by Threatlint… + issue link"],
    ["Closed state", "Closed / False Positive (from sys_choice)", "Closed (Completed category of the type)", "dismissed · false positive"],
    ["CodeSecurity rule", "source/scanner/category/type field", "Tags, title, or area path", "Code-scanning alerts only"],
    ["Credential", "OAuth client via AWS SSO (private window)", "PAT: Work Items read & write", "Token: code scanning alerts + issues write"],
], Inches(0.6), Inches(1.8), [Inches(2.2), Inches(3.4), Inches(3.3), Inches(3.3)], size=12, row_h=0.6)

# 5. Efficacy method
slide = light("How efficacy was measured", "Evaluation method")
fp_rows = [row for row in rows if row["expected"] == "FALSE POSITIVE"]
card(slide, Inches(0.6), Inches(1.8), Inches(6), Inches(4.6), "Inputs",
     f"• {m['total']} CodeSecurity AVITs produced by real Semgrep registry rules on an evaluation repository\n"
     f"• Ground truth set by reading the code: {len(fp_rows)} false positives, {m['tpTotal']} true positives\n"
     f"• 1 non-CodeSecurity AVIT (DAST) mixed in to test the scope filter\n• Mode: {m['mode']}\n• Model: {m['model']}", size=14)
card(slide, Inches(6.8), Inches(1.8), Inches(5.95), Inches(4.6), "What was scored",
     "• Disposition per AVIT vs. ground truth\n• Unsafe closures: real vulnerabilities marked FALSE POSITIVE (the costly error)\n"
     "• Scope: non-CodeSecurity AVITs excluded before triage\n• Closure mechanics: attach-before-close, AVIT-prefixed files, TRUE POSITIVE refused\n"
     "Closing ran against recording stand-ins; no real tracker was changed.", accent=NAVY, size=14)

# 6. Results headline
slide = light("Results", "Efficacy")
stats = [(f"{m['correct']}/{m['total']}", "dispositions correct", NAVY),
         (f"{m['fpCorrect']}/{m['fpTotal']}", "false positives identified", GREEN),
         (f"{m['tpCorrect']}/{m['tpTotal']}", "true positives kept open", GREEN),
         (str(m["unsafeClosures"]), "unsafe closures", GREEN if m["unsafeClosures"] == 0 else RED)]
for index, (value, label, color) in enumerate(stats):
    x = Inches(0.6 + index * 3.1)
    rect(slide, x, Inches(1.9), Inches(2.9), Inches(2.2), WHITE, LINE)
    txt(slide, value, x, Inches(2.1), Inches(2.9), Inches(1.1), 48, True, color, PP_ALIGN.CENTER)
    txt(slide, label, x, Inches(3.3), Inches(2.9), Inches(0.6), 15, False, GREY, PP_ALIGN.CENTER)
excluded = ", ".join(m["excludedNonCodeSecurity"]) or "none"
txt(slide, f"Accuracy {pct(m['correct'], m['total'])} · ambiguous {m['ambiguous']} · non-CodeSecurity excluded: {excluded} · "
           f"{m['toolCalls']} repository inspections · {m['seconds']} s end to end",
    Inches(0.6), Inches(4.5), Inches(12.2), Inches(0.6), 15, False, INK)
txt(slide, "Small sample: this shows the workflow is sound, not a production accuracy rate. Re-run on live AVITs with ground truth (docs/fp-closure.md).",
    Inches(0.6), Inches(5.2), Inches(12.2), Inches(0.8), 13, False, GREY)

# 7. Per-AVIT verdicts
slide = light("Verdict for each AVIT", "Efficacy detail")
detail = [["AVIT", "Rule (short)", "Location", "Expected", "Agent", "✓"]]
for row in rows:
    detail.append([row["avit"], row.get("rule", "").split(".")[-1], row.get("location", ""), row["expected"], f"{row['disposition']} ({row.get('confidence', '?')})", "✓" if row["correct"] else "✗"])
table(slide, detail, Inches(0.6), Inches(1.7), [Inches(1.5), Inches(3.4), Inches(1.9), Inches(1.9), Inches(2.8), Inches(0.6)], size=11, row_h=0.5)
reasons = "\n".join(f"{row['avit']}: {row['why']}" for row in rows if row.get("why"))
txt(slide, reasons, Inches(0.6), Inches(5.4), Inches(12.2), Inches(1.9), 11, False, GREY)

# 8. Closure evidence
slide = light("Closure: evidence first, then state", "Closure mechanics")
evidence = [["AVIT", "Outcome", "Report attached", "ServiceNow calls (in order)"]]
for item in closures:
    if item.get("refused"):
        evidence.append([item["avit"], f"Refused (HTTP {item['status']})", "—", item["refused"][:70]])
    else:
        calls = [call for call in item["calls"] if "oauth" not in call]
        short = {"table/sn_vul_app_vulnerable_item": "read AVIT", "table/sys_choice": "read choices", "attachment/file": "attach report"}
        names = [next((label for key, label in short.items() if key in call and not call.startswith("PATCH")), "close (PATCH)" if call.startswith("PATCH") else call) for call in calls]
        evidence.append([item["avit"], "Closed / False Positive" if item["closed"] else "Not closed", ", ".join(item["attachments"]), " → ".join(names)])
table(slide, evidence, Inches(0.6), Inches(1.7), [Inches(1.5), Inches(2.3), Inches(3.6), Inches(4.8)], size=10, row_h=0.55)
txt(slide, "Each closure uploads the AVIT-prefixed report before the PATCH that closes it. A TRUE POSITIVE is refused by the server even if the UI is bypassed.",
    Inches(0.6), Inches(5.8), Inches(12.2), Inches(0.8), 14, False, INK)

# 9. Guardrails
slide = light("Guardrails", "Safety")
card(slide, Inches(0.6), Inches(1.8), Inches(3.9), Inches(4.6), "Human in the loop",
     "Every close and every Semgrep ignore asks for confirmation. A finding must be FALSE POSITIVE and reference exactly one item.")
card(slide, Inches(4.72), Inches(1.8), Inches(3.9), Inches(4.6), "Narrow scope",
     "Only CodeSecurity items qualify. Semgrep ignores are limited to one rule, project, and file, and the server re-matches before triaging.", accent=NAVY)
card(slide, Inches(8.84), Inches(1.8), Inches(3.9), Inches(4.6), "Credentials stay put",
     "Tokens live in session memory only and are never written to disk. Redirects are refused, so tokens go only to validated API hosts.", accent=GREEN)

# 10. Next steps
slide = light("Next steps", "Rollout")
card(slide, Inches(0.6), Inches(1.8), Inches(12.15), Inches(4.6), "To run on production items",
     "1. Grant the ServiceNow OAuth client write access to sn_vul_app_vulnerable_item and sys_attachment\n"
     "2. Issue ADO PATs (Work Items: read & write) and GitHub tokens (code scanning alerts + issues: write)\n"
     "3. Confirm how AVR labels the CodeSecurity view; set THREATLINT_SNOW_CODE_SECURITY_PATTERN if needed\n"
     "4. Run the efficacy evaluation on 20–50 live CodeSecurity AVITs with reviewer-labelled ground truth\n"
     "5. Route rule-logic changes (pattern-not, path exclusions) from FP reports to the Semgrep rule owner", size=16)

prs.save(OUT)
print(f"Wrote {OUT} ({len(prs.slides)} slides)")
